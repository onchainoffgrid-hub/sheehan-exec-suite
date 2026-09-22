#!/usr/bin/env python3
"""Build unified BDR TAM dialer JSON from real on-disk lists.
Never invent emails/phones. Dedupe: email → phone → org+name.
Default rank: Tier1 first within sector, then has_email, then has_phone, then name.
"""
from __future__ import annotations

import csv
import json
import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

try:
    import openpyxl
except ImportError:
    import subprocess
    subprocess.check_call(["pip", "install", "openpyxl", "-q"])
    import openpyxl

ET = ZoneInfo("America/New_York")
OUT = Path("/workspace/dashboard/exec-suite/data/bdr_tam_dialer.json")
AS_OF = datetime.now(ET).strftime("%Y-%m-%d %H:%M ET")

# Canonical segment labels (chip order — Tier1 focus sectors first)
SEGMENT_ORDER = [
    "Senior / IL-AL",
    "PPEC",
    "Adult Daycare",
    "CDD / HOA",
    "Luxury Apartment",
    "Golf / Country Club",
    "Corporate HR",
    "Corporate QSR / Field",
    "Industrial / Employee Apprec",
    "University",
    "Independent School",
    "VPK / ELC",
    "Daycare / Childcare",
    "Goddard / Fancy School",
    "PTA / PTO",
    "Church / Faith",
    "Festival / Events",
    "Brewery / Family",
    "Car Dealership",
    "Hotel / Kids Club",
    "Preservation / SPAR",
    "Nonprofit",
    "Hospital / Red Tape",
    "Municipal / B2G",
    "B2C / Violet / Inbound",
    "Existing / Customer",
    "Other",
]

SEG_ALIASES = {
    # senior
    "elite il/al / senior campuses": "Senior / IL-AL",
    "assisted living / il-al": "Senior / IL-AL",
    "sr_memory_al": "Senior / IL-AL",
    "senior": "Senior / IL-AL",
    "senior care": "Senior / IL-AL",
    "memory care": "Senior / IL-AL",
    "memory": "Senior / IL-AL",
    "senior center (gov)": "Municipal / B2G",
    "senior center (nonprofit)": "Senior / IL-AL",
    "sr_nonprofit_coa": "Senior / IL-AL",
    "sr_gov_b2g": "Municipal / B2G",
    "al / memory": "Senior / IL-AL",
    "assisted living": "Senior / IL-AL",
    # ppec / adc
    "ppec": "PPEC",
    "adult_daycare": "Adult Daycare",
    "adult day care": "Adult Daycare",
    "adult daycare": "Adult Daycare",
    # cdd
    "cdd/hoa": "CDD / HOA",
    "cdd_hoa": "CDD / HOA",
    "cdd/hoa_vesta": "CDD / HOA",
    "hoa": "CDD / HOA",
    "cdd": "CDD / HOA",
    # luxury / golf
    "luxury_apartment": "Luxury Apartment",
    "luxury apartment": "Luxury Apartment",
    "golf_country_club": "Golf / Country Club",
    "golf / country club": "Golf / Country Club",
    # corp
    "corporate_hr": "Corporate HR",
    "corporate-hr": "Corporate HR",
    "corp hr": "Corporate HR",
    "corporate_qsr_field": "Corporate QSR / Field",
    "qsr-field": "Corporate QSR / Field",
    "corporate_qsr": "Corporate QSR / Field",
    "field_corporate": "Corporate QSR / Field",
    "corporate": "Corporate HR",
    "enterprise": "Corporate HR",
    "regional_industrial_employeeapprec": "Industrial / Employee Apprec",
    # edu
    "university": "University",
    "independent_school": "Independent School",
    "independent school": "Independent School",
    "indie school": "Independent School",
    "school": "Independent School",
    "goddard_premium": "Goddard / Fancy School",
    "fancy school": "Goddard / Fancy School",
    "vpk": "VPK / ELC",
    "vpk_elc": "VPK / ELC",
    "daycare_chain": "Daycare / Childcare",
    "daycare": "Daycare / Childcare",
    "childcare/daycare": "Daycare / Childcare",
    "faith school": "Church / Faith",
    # pta church
    "pta": "PTA / PTO",
    "church": "Church / Faith",
    # fest / brewery
    "fest_events": "Festival / Events",
    "community event": "Festival / Events",
    "corporate event": "Festival / Events",
    "brewery_family": "Brewery / Family",
    "car_dealership": "Car Dealership",
    "hotel_resort_kidsclub": "Hotel / Kids Club",
    "hotel": "Hotel / Kids Club",
    # other
    "preservation": "Preservation / SPAR",
    "nonprofit": "Nonprofit",
    "hospital_insurance": "Hospital / Red Tape",
    "hospital/insurance": "Hospital / Red Tape",
    "hospital": "Hospital / Red Tape",
    "red_tape": "Hospital / Red Tape",
    "municipal parks": "Municipal / B2G",
    "b2b hubspot": "Other",
    "recurring_weekday": "Senior / IL-AL",
    "violet": "B2C / Violet / Inbound",
    "existing": "Existing / Customer",
    "sponsor_retail": "Festival / Events",
    "other": "Other",
    "homeschool": "Independent School",
    "priority calls": "Festival / Events",
    "04_bant_recurring": "Senior / IL-AL",
    "bant recurring": "Senior / IL-AL",
    "recurring": "Senior / IL-AL",
}


def norm_seg(raw: str | None) -> str:
    if not raw:
        return "Other"
    s = str(raw).strip()
    key = s.lower().replace("–", "-").replace("—", "-")
    if key in SEG_ALIASES:
        return SEG_ALIASES[key]
    # partial
    for a, canon in SEG_ALIASES.items():
        if a in key or key in a:
            return canon
    # title-ish fallbacks
    if "senior" in key or "memory" in key or "assisted" in key or "il/al" in key or "elite" in key:
        return "Senior / IL-AL"
    if "cdd" in key or "hoa" in key or "vesta" in key:
        return "CDD / HOA"
    if "ppec" in key:
        return "PPEC"
    if "vpk" in key:
        return "VPK / ELC"
    if "brew" in key:
        return "Brewery / Family"
    if "church" in key or "faith" in key:
        return "Church / Faith"
    if "pta" in key or "pto" in key:
        return "PTA / PTO"
    if "golf" in key:
        return "Golf / Country Club"
    if "university" in key or "college" in key:
        return "University"
    if "hospital" in key or "insurance" in key or "red tape" in key:
        return "Hospital / Red Tape"
    if "goddard" in key or "bolles" in key or "episcopal" in key:
        return "Goddard / Fancy School"
    if "inbound" in key or "violet" in key or "b2c" in key:
        return "B2C / Violet / Inbound"
    if "corp" in key and ("qsr" in key or "field" in key or "mcdonald" in key):
        return "Corporate QSR / Field"
    if "corp" in key or "hr" in key:
        return "Corporate HR"
    return s if len(s) < 40 else "Other"


def clean_email(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    if not s or s.lower() in ("none", "null", "n/a", "na", "-", "—"):
        return ""
    # strip extras
    s = s.split(";")[0].split(",")[0].strip()
    if "@" not in s or " " in s.split("@")[0]:
        # try extract
        m = re.search(r"[\w.+-]+@[\w.-]+\.\w+", s)
        if m:
            s = m.group(0)
        else:
            return ""
    return s.lower()


def clean_phone(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    if not s or s.lower() in ("none", "null", "n/a", "na", "-", "—"):
        return ""
    digits = re.sub(r"\D", "", s)
    if len(digits) < 7:
        return ""
    # keep original formatting if looks phone-like
    return s


def phone_key(p: str) -> str:
    d = re.sub(r"\D", "", p or "")
    if len(d) == 11 and d.startswith("1"):
        d = d[1:]
    return d


def clean_name(*parts) -> str:
    bits = []
    for p in parts:
        if p is None:
            continue
        s = str(p).strip()
        if s and s.lower() not in ("none", "null", "n/a", "nan"):
            bits.append(s)
    return " ".join(bits).strip()


def is_tier1(call_band: str, tier: str, s_tier: str, source: str, extra: dict | None = None) -> bool:
    cb = (call_band or "").strip().lower()
    t = (tier or "").strip().lower()
    st = (s_tier or "").strip().upper()
    if st in ("S", "TIER1", "TIER 1", "1"):
        return True
    if "call now" in cb:
        return True
    if cb in ("call now", "call_now", "call"):
        return True
    if t.startswith("elite") or t in ("elite_il", "elite_al", "elite_top25"):
        return True
    if "tier 1" in t or t == "tier1" or t == "1":
        return True
    # WAVE / fall priority often Tier1-ish
    if source.startswith("wave_c") and st == "S":
        return True
    if source.startswith("elite"):
        return True
    if source.startswith("inbound_fire") and (extra or {}).get("priority") in (1, "1", 2, "2"):
        return True
    if source.startswith("enterprise_call_now") or source.endswith("CALL_NOW"):
        return True
    return False


def next_action(row: dict) -> str:
    if row.get("is_tier1") and row.get("email"):
        return "Email + call (Tier 1)"
    if row.get("is_tier1") and row.get("phone"):
        return "Call now"
    if row.get("email"):
        return "Email intro"
    if row.get("phone"):
        return "Cold call / tap"
    if row.get("form_url"):
        return "Submit contact form"
    return "Research contact"


def make_row(**kw) -> dict | None:
    account = (kw.get("account") or "").strip()
    if not account:
        return None
    email = clean_email(kw.get("email"))
    phone = clean_phone(kw.get("phone"))
    contact = clean_name(kw.get("contact"), kw.get("contact_first"), kw.get("contact_last"))
    # Prefer explicit contact over first+last already merged
    if kw.get("contact"):
        contact = clean_name(kw.get("contact"))
    elif kw.get("contact_first") or kw.get("contact_last"):
        contact = clean_name(kw.get("contact_first"), kw.get("contact_last"))
    title = (kw.get("title") or "").strip()
    segment = norm_seg(kw.get("segment"))
    raw_seg = (kw.get("segment") or "").strip()
    tier = (kw.get("tier") or "").strip()
    call_band = (kw.get("call_band") or "").strip()
    s_tier = (kw.get("s_tier") or "").strip()
    penetration = (kw.get("penetration") or kw.get("status") or "").strip()
    source = (kw.get("source") or "").strip()
    form_url = (kw.get("form_url") or "").strip()
    county = (kw.get("county") or "").strip()
    city = (kw.get("city") or "").strip()
    notes = (kw.get("notes") or "").strip()
    # Skip empty shells with zero info? Keep for TAM coverage but flag
    has_info = bool(email or phone or contact or form_url)
    t1 = is_tier1(call_band, tier, s_tier, source, kw)
    # Derive display tier
    if t1:
        tier_display = tier or "Tier 1"
        if not tier and "call now" in call_band.lower():
            tier_display = "Tier 1 / Call now"
        elif not tier and s_tier.upper() == "S":
            tier_display = "Tier 1 / S"
    else:
        tier_display = tier or call_band or s_tier or ""
    row = {
        "segment": segment,
        "segment_raw": raw_seg,
        "account": account,
        "contact": contact,
        "title": title,
        "phone": phone,
        "email": email,
        "tier": tier_display,
        "is_tier1": t1,
        "call_band": call_band,
        "penetration": penetration,
        "next_action": "",  # fill after
        "source": source,
        "county": county,
        "city": city,
        "form_url": form_url,
        "notes": notes[:200] if notes else "",
        "has_email": bool(email),
        "has_phone": bool(phone),
        "has_info": has_info,
    }
    row["next_action"] = next_action(row)
    return row


def richness(r: dict) -> int:
    return (
        (4 if r.get("email") else 0)
        + (2 if r.get("phone") else 0)
        + (1 if r.get("contact") else 0)
        + (1 if r.get("title") else 0)
        + (1 if r.get("is_tier1") else 0)
        + (1 if r.get("form_url") else 0)
    )


# Segments that must not be overwritten by a broader sibling label
SPECIFIC_SEGS = {
    "PPEC", "Adult Daycare", "Goddard / Fancy School", "Brewery / Family",
    "Car Dealership", "Hotel / Kids Club", "Luxury Apartment", "Golf / Country Club",
    "VPK / ELC", "Preservation / SPAR", "Industrial / Employee Apprec",
    "Existing / Customer", "B2C / Violet / Inbound",
}
WEAK_SEGS = {"Other", ""}


def pick_segment(a: str, b: str) -> str:
    a = a or "Other"
    b = b or "Other"
    if a in WEAK_SEGS and b not in WEAK_SEGS:
        return b
    if b in WEAK_SEGS:
        return a
    if a in SPECIFIC_SEGS:
        return a
    if b in SPECIFIC_SEGS:
        return b
    return a  # keep primary


def merge_rows(a: dict, b: dict) -> dict:
    # Prefer richer; fill blanks from other. Never drop a real email — note alt.
    primary, secondary = (a, b) if richness(a) >= richness(b) else (b, a)
    out = dict(primary)
    for k, v in secondary.items():
        if k in ("is_tier1", "has_email", "has_phone", "has_info", "email", "segment", "segment_raw"):
            continue
        if not out.get(k) and v:
            out[k] = v
    # segment: keep specific labels; only replace weak Other
    chosen = pick_segment(out.get("segment") or "", secondary.get("segment") or "")
    out["segment"] = chosen
    if chosen == (secondary.get("segment") or "") and secondary.get("segment_raw"):
        out["segment_raw"] = secondary["segment_raw"]
    # emails: keep primary; stash alternate in notes if different
    pe, se = (primary.get("email") or ""), (secondary.get("email") or "")
    if pe and se and pe != se:
        note = out.get("notes") or ""
        alt = f"alt_email:{se}"
        if alt not in note:
            out["notes"] = (note + " · " + alt).strip(" ·")
    elif not pe and se:
        out["email"] = se
    # union sources
    sa = set(filter(None, (primary.get("source") or "").split("|")))
    sb = set(filter(None, (secondary.get("source") or "").split("|")))
    out["source"] = "|".join(sorted(sa | sb))
    out["is_tier1"] = bool(primary.get("is_tier1") or secondary.get("is_tier1"))
    out["has_email"] = bool(out.get("email"))
    out["has_phone"] = bool(out.get("phone"))
    out["has_info"] = bool(out.get("email") or out.get("phone") or out.get("contact") or out.get("form_url"))
    if out["is_tier1"] and "Tier 1" not in (out.get("tier") or ""):
        out["tier"] = (out.get("tier") + " · Tier 1").strip(" ·") if out.get("tier") else "Tier 1"
    out["next_action"] = next_action(out)
    return out


class Dialer:
    def __init__(self):
        self.by_email: dict[str, dict] = {}
        self.by_phone: dict[str, dict] = {}
        self.by_org_name: dict[str, dict] = {}
        self.rows: list[dict] = []
        self.source_counts: Counter = Counter()
        self.skipped_empty = 0

    def _org_key(self, account: str, contact: str) -> str:
        a = re.sub(r"[^a-z0-9]", "", (account or "").lower())
        c = re.sub(r"[^a-z0-9]", "", (contact or "").lower())
        return f"{a}::{c}" if c else a

    def add(self, row: dict | None):
        if not row:
            return
        email = row.get("email") or ""
        phone = row.get("phone") or ""
        pk = phone_key(phone)
        ok = self._org_key(row["account"], row.get("contact") or "")

        existing = None
        # 1) Same email → merge
        if email and email in self.by_email:
            existing = self.by_email[email]
        else:
            # 2) Same phone → merge ONLY if emails don't conflict
            cand = self.by_phone.get(pk) if pk else None
            if cand is not None:
                ce = cand.get("email") or ""
                if not email or not ce or email == ce:
                    existing = cand
            # 3) Same org+name → merge ONLY if emails don't conflict
            if existing is None and ok and ok in self.by_org_name:
                cand = self.by_org_name[ok]
                ce = cand.get("email") or ""
                if not email or not ce or email == ce:
                    existing = cand

        if existing:
            merged = merge_rows(existing, row)
            try:
                idx = self.rows.index(existing)
                self.rows[idx] = merged
            except ValueError:
                self.rows.append(merged)
            # drop stale index pointers to old object
            for d in (self.by_email, self.by_phone, self.by_org_name):
                for k, v in list(d.items()):
                    if v is existing:
                        d[k] = merged
            if merged.get("email"):
                self.by_email[merged["email"]] = merged
            if phone_key(merged.get("phone") or ""):
                self.by_phone[phone_key(merged["phone"])] = merged
            self.by_org_name[self._org_key(merged["account"], merged.get("contact") or "")] = merged
            # also index the other email if noted — keep as separate row instead
            self.source_counts[row.get("source", "?")] += 1
            # If emails conflicted we did not merge above; but if merge kept alt in notes
            # and row had a distinct email, also emit a second contact row.
            se = row.get("email") or ""
            pe = merged.get("email") or ""
            if se and pe and se != pe:
                # clone secondary-email contact as its own dial row
                alt = dict(row)
                alt["notes"] = ((row.get("notes") or "") + f" · same_account_as:{merged.get('account')}").strip(" ·")
                self._index_new(alt)
            return

        self._index_new(row)
        self.source_counts[row.get("source", "?")] += 1

    def _index_new(self, row: dict):
        self.rows.append(row)
        email = row.get("email") or ""
        pk = phone_key(row.get("phone") or "")
        ok = self._org_key(row["account"], row.get("contact") or "")
        if email:
            self.by_email[email] = row
        # Only claim phone index if free or same email
        if pk and (pk not in self.by_phone or (self.by_phone[pk].get("email") or "") in ("", email)):
            self.by_phone[pk] = row
        if ok and (ok not in self.by_org_name or (self.by_org_name[ok].get("email") or "") in ("", email)):
            self.by_org_name[ok] = row


def main():
    D = Dialer()

    # ---------- Sources ----------

    def load_json_list(path: str):
        p = Path(path)
        if not p.exists():
            return []
        data = json.loads(p.read_text())
        return data if isinstance(data, list) else []


    # 1) Enterprise / BDR primary dialers
    for path, src in [
        ("/workspace/dashboard/exec-suite/data/enterprise_call_now.json", "enterprise_call_now"),
        ("/workspace/dashboard/exec-suite/data/enterprise_dialer.json", "enterprise_dialer"),
        ("/workspace/dashboard/exec-suite/data/bdr_enterprise_dialer.json", "bdr_enterprise_dialer"),
        ("/workspace/dashboard/exec-suite/data/bdr_primary_dialer.json", "bdr_primary_dialer"),
    ]:
        for r in load_json_list(path):
            D.add(
                make_row(
                    account=r.get("company_name"),
                    contact_first=r.get("contact_first"),
                    contact_last=r.get("contact_last"),
                    title=r.get("title_persona"),
                    phone=r.get("phone"),
                    email=r.get("email"),
                    segment=r.get("segment"),
                    call_band=r.get("call_band"),
                    county=r.get("county"),
                    city=r.get("city"),
                    notes=r.get("notes"),
                    source=src,
                    penetration=r.get("buyer_motion") or "",
                )
            )

    # 2) Fall priority
    for r in load_json_list("/workspace/dashboard/exec-suite/data/fall_priority.json"):
        D.add(
            make_row(
                account=r.get("company") or r.get("event"),
                contact=r.get("name"),
                phone=r.get("phone"),
                email=r.get("email"),
                segment=r.get("segment") or "Festival / Events",
                call_band="Call now",  # fall heat = act now
                county=r.get("county"),
                notes=r.get("why"),
                source="fall_priority",
                penetration="Fall heat",
                tier="Tier 1 / Fall",
            )
        )

    # 3) Org whitespace contacts
    ow_path = Path("/workspace/dashboard/exec-suite/data/org_whitespace.json")
    if ow_path.exists():
        ow = json.loads(ow_path.read_text())
        for o in ow.get("orgs") or []:
            contacts = o.get("contacts") or []
            if not contacts:
                D.add(
                    make_row(
                        account=o.get("account"),
                        segment=o.get("segment"),
                        tier=o.get("tier"),
                        call_band=o.get("call_band"),
                        penetration=o.get("penetration") or o.get("account_status"),
                        city=o.get("city"),
                        form_url=o.get("contact_form_url"),
                        source="org_whitespace",
                    )
                )
            for c in contacts:
                D.add(
                    make_row(
                        account=o.get("account"),
                        contact=c.get("name"),
                        title=c.get("title"),
                        phone=c.get("phone"),
                        email=c.get("email"),
                        segment=o.get("segment"),
                        tier=o.get("tier"),
                        call_band=o.get("call_band"),
                        penetration=o.get("penetration") or o.get("account_status"),
                        city=o.get("city"),
                        form_url=c.get("form_url") or o.get("contact_form_url"),
                        source="org_whitespace",
                    )
                )

    # 4) Elite Senior xlsx
    elite = Path("/workspace/bdr/lists/Elite_Senior_Contact_Forms.xlsx")
    if elite.exists():
        wb = openpyxl.load_workbook(elite, read_only=True, data_only=True)
        for sheet in ("Elite_Top25", "Penetration_All_Seniors", "Phone_Tap_Queue"):
            if sheet not in wb.sheetnames:
                continue
            ws = wb[sheet]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            hdr = [str(h or "").strip() for h in rows[0]]

            def col(*names):
                for n in names:
                    if n in hdr:
                        return hdr.index(n)
                return None

            for r in rows[1:]:
                if not r:
                    continue
                def g(n):
                    i = col(n) if isinstance(n, str) else None
                    # multi
                    return None

                def get(*names):
                    for n in names:
                        if n in hdr:
                            v = r[hdr.index(n)]
                            if v is not None and str(v).strip():
                                return v
                    return None

                prop = get("property_name")
                if not prop:
                    continue
                # Parse activities_contact for name/email hints (do not invent)
                act = str(get("activities_contact") or "")
                contact = ""
                # "Patricia Valenzuela, Director..." → name before comma if no email-only
                if act and "@" not in act.split(",")[0]:
                    contact = act.split(";")[0].split(",")[0].strip()
                    if len(contact) > 60:
                        contact = ""
                D.add(
                    make_row(
                        account=prop,
                        contact=contact,
                        phone=get("phone"),
                        email=get("email_if_known", "email"),
                        segment="Senior / IL-AL",
                        tier=get("tier") or ("Elite_Top25" if sheet == "Elite_Top25" else ""),
                        call_band="Call now" if sheet in ("Elite_Top25", "Phone_Tap_Queue") else "",
                        penetration=get("penetration_status"),
                        city=get("city"),
                        form_url=get("contact_form_url"),
                        notes=get("notes") or get("form_fill_notes"),
                        source=f"elite_senior:{sheet}",
                        s_tier="S" if sheet == "Elite_Top25" else "",
                    )
                )
        wb.close()

    # 5) INBOUND_FIRE
    inf = Path("/workspace/bdr/lists/INBOUND_FIRE.xlsx")
    if inf.exists():
        wb = openpyxl.load_workbook(inf, read_only=True, data_only=True)
        if "Action_Queue" in wb.sheetnames:
            ws = wb["Action_Queue"]
            rows = list(ws.iter_rows(values_only=True))
            hdr = [str(h or "").strip() for h in rows[0]]
            for r in rows[1:]:
                if not r:
                    continue
                def get(n):
                    return r[hdr.index(n)] if n in hdr else None
                name = get("name")
                email = get("email")
                if not name and not email:
                    continue
                # account = name org part if " / " present
                account = str(name or email or "").split("/")[-1].strip() if name else str(email)
                contact = str(name).split("/")[0].strip() if name and "/" in str(name) else str(name or "")
                D.add(
                    make_row(
                        account=account or contact,
                        contact=contact,
                        email=email,
                        phone=get("phone"),
                        segment="B2C / Violet / Inbound",
                        call_band="Call now",
                        notes=get("why"),
                        source="inbound_fire:Action_Queue",
                        priority=get("priority"),
                        penetration=str(get("kind") or ""),
                        tier="Tier 1 / Inbound",
                    )
                )
        if "GoDaddy_Inquiries" in wb.sheetnames:
            ws = wb["GoDaddy_Inquiries"]
            rows = list(ws.iter_rows(values_only=True))
            hdr = [str(h or "").strip() for h in rows[0]]
            for r in rows[1:]:
                if not r:
                    continue
                def get(n):
                    return r[hdr.index(n)] if n in hdr else None
                name = get("name")
                email = get("email")
                if not name and not email:
                    continue
                D.add(
                    make_row(
                        account=str(name or email),
                        contact=str(name or ""),
                        email=email,
                        phone=get("phone"),
                        segment="B2C / Violet / Inbound",
                        call_band="Call now" if str(get("status") or "") == "NEEDS_REPLY" else "Warm",
                        notes=get("action_note"),
                        source="inbound_fire:GoDaddy",
                        penetration=str(get("status") or ""),
                    )
                )
        wb.close()

    # 6) WAVE_C email (high quality Tier1)
    wave = Path("/workspace/revops/WAVE_C_CALL_NOW_email.csv")
    if wave.exists():
        with wave.open(newline="", encoding="utf-8", errors="replace") as f:
            for r in csv.DictReader(f):
                D.add(
                    make_row(
                        account=r.get("company_name"),
                        contact=r.get("contact"),
                        title=r.get("title"),
                        phone=r.get("phone"),
                        email=r.get("email"),
                        segment=r.get("sector") or r.get("seat"),
                        call_band=r.get("call_band"),
                        s_tier=r.get("s_tier"),
                        county=r.get("county"),
                        penetration=r.get("penetration"),
                        notes=r.get("one_liner"),
                        source="wave_c_call_now_email",
                        tier="Tier 1" if (r.get("s_tier") or "").upper() == "S" else r.get("s_tier"),
                    )
                )

    # 7) RevOps CSVs with contacts
    CSV_MAP = [
        ("/workspace/revops/ppec.csv", "revops:ppec", "PPEC"),
        ("/workspace/revops/senior_centers.csv", "revops:senior_centers", "Senior / IL-AL"),
        ("/workspace/revops/al_memory.csv", "revops:al_memory", "Senior / IL-AL"),
        ("/workspace/revops/adult_daycare.csv", "revops:adult_daycare", "Adult Daycare"),
        ("/workspace/revops/corporate.csv", "revops:corporate", None),
        ("/workspace/revops/universities.csv", "revops:universities", "University"),
        ("/workspace/revops/preservation.csv", "revops:preservation", "Preservation / SPAR"),
        ("/workspace/revops/pta_officers.csv", "revops:pta_officers", "PTA / PTO"),
        ("/workspace/revops/b2b_expand_cdd.csv", "revops:b2b_expand_cdd", "CDD / HOA"),
        ("/workspace/revops/b2b_expand_vpk.csv", "revops:b2b_expand_vpk", "VPK / ELC"),
        ("/workspace/revops/b2b_expand_churches.csv", "revops:b2b_expand_churches", "Church / Faith"),
        ("/workspace/revops/b2b_expand_golf_clubs.csv", "revops:b2b_expand_golf_clubs", "Golf / Country Club"),
        ("/workspace/revops/b2b_expand_indie_schools.csv", "revops:b2b_expand_indie_schools", "Independent School"),
        ("/workspace/revops/b2b_expand_luxury_apartments.csv", "revops:b2b_expand_luxury_apt", "Luxury Apartment"),
        ("/workspace/revops/Jax_Corporate_BDR_CallList.csv", "revops:corp_bdr_calllist", None),
        ("/workspace/revops/CALL_NOW_phone.csv", "revops:call_now_phone", None),
        ("/workspace/revops/DIALER_phone.csv", "revops:dialer_phone", None),
    ]
    for path, src, force_seg in CSV_MAP:
        p = Path(path)
        if not p.exists():
            continue
        with p.open(newline="", encoding="utf-8", errors="replace") as f:
            reader = csv.DictReader(f)
            for r in reader:
                # flexible field names
                account = (
                    r.get("company_name")
                    or r.get("company")
                    or r.get("Company Name")
                    or r.get("school_name")
                    or ""
                )
                contact = (
                    r.get("contact")
                    or clean_name(r.get("first_name"), r.get("last_name"))
                    or clean_name(r.get("ContactFirst"), r.get("ContactLast"))
                    or ""
                )
                email = r.get("email") or r.get("Email") or ""
                phone = r.get("phone") or r.get("Phone Number") or r.get("mobile_phone") or ""
                title = r.get("title") or r.get("title_persona") or r.get("persona") or r.get("Job Title") or ""
                seg = force_seg or r.get("segment") or r.get("seat") or r.get("tab") or ""
                call_band = r.get("call_band") or r.get("Call Band") or ""
                if src.endswith("call_now_phone") or "CALL_NOW" in (r.get("call_band") or ""):
                    call_band = call_band or "Call now"
                if r.get("tier") and str(r.get("tier")).upper() in ("S", "A", "TIER1", "1"):
                    pass
                D.add(
                    make_row(
                        account=account,
                        contact=contact,
                        title=title,
                        phone=phone,
                        email=email,
                        segment=seg,
                        call_band=call_band or ("Call now" if "call_now" in src else ""),
                        tier=r.get("tier") or r.get("s_tier") or r.get("market_tier") or "",
                        s_tier=r.get("s_tier") or "",
                        county=r.get("county") or r.get("County"),
                        city=r.get("city"),
                        penetration=r.get("penetration") or r.get("decision_speed") or "",
                        notes=r.get("one_liner") or r.get("notes") or "",
                        source=src,
                    )
                )

    # 8) BANT HubSpot contacts + CALL_NOW_ALL (quality)
    bant = Path("/workspace/revops/Jax_BANT_Dialer_TAM.xlsx")
    if bant.exists():
        wb = openpyxl.load_workbook(bant, read_only=True, data_only=True)
        # Contacts_HubSpot
        if "Contacts_HubSpot" in wb.sheetnames:
            ws = wb["Contacts_HubSpot"]
            rows = list(ws.iter_rows(values_only=True))
            hdr = [str(h or "").strip() for h in rows[0]]
            for r in rows[1:]:
                if not r:
                    continue
                def get(n):
                    return r[hdr.index(n)] if n in hdr else None
                email = get("Email")
                company = get("Company Name")
                if not company and not email:
                    continue
                D.add(
                    make_row(
                        account=company or email,
                        contact=clean_name(get("First Name"), get("Last Name")),
                        title=get("Job Title"),
                        phone=get("Phone Number"),
                        email=email,
                        segment=get("Segment"),
                        call_band=get("Call Band"),
                        tier=get("Market Tier"),
                        county=get("County"),
                        penetration=get("Status"),
                        source="bant:Contacts_HubSpot",
                    )
                )
        # 11_CALL_NOW_ALL — Tier 1 across BANT
        if "11_CALL_NOW_ALL" in wb.sheetnames:
            ws = wb["11_CALL_NOW_ALL"]
            rows = list(ws.iter_rows(values_only=True))
            # row0 is title banner, row1 is header
            hdr = None
            start = 0
            for i, r in enumerate(rows[:5]):
                if r and "company_name" in [str(x or "").lower() for x in r]:
                    hdr = [str(h or "").strip() for h in r]
                    start = i + 1
                    break
            if hdr:
                for r in rows[start:]:
                    if not r:
                        continue
                    def get(n):
                        nl = n.lower()
                        for i, h in enumerate(hdr):
                            if h.lower() == nl:
                                return r[i] if i < len(r) else None
                        return None
                    company = get("company_name")
                    if not company:
                        continue
                    D.add(
                        make_row(
                            account=company,
                            contact=get("contact"),
                            title=get("title"),
                            phone=get("phone"),
                            email=get("email"),
                            segment=get("tab_label") or get("market_tier"),
                            call_band="Call now",
                            tier=get("bant_tier") or "Tier 1",
                            county=get("county"),
                            penetration=get("penetration"),
                            source="bant:11_CALL_NOW_ALL",
                            s_tier="S",
                        )
                    )
        # Recurring seniors tab for subscription Tier1
        for sheet, force in [
            ("04_bant_Recurring", "Senior / IL-AL"),
            ("03_FANCY_SCHOOLS", "Goddard / Fancy School"),
        ]:
            if sheet not in wb.sheetnames:
                continue
            ws = wb[sheet]
            rows = list(ws.iter_rows(values_only=True))
            hdr = None
            start = 0
            for i, r in enumerate(rows[:5]):
                if r and any(str(x or "").lower() == "company_name" for x in r):
                    hdr = [str(h or "").strip() for h in r]
                    start = i + 1
                    break
            if not hdr:
                continue
            for r in rows[start:]:
                if not r:
                    continue
                def get(n):
                    for i, h in enumerate(hdr):
                        if h.lower() == n.lower():
                            return r[i] if i < len(r) else None
                    return None
                company = get("company_name")
                if not company:
                    continue
                bant_tier = str(get("bant_tier") or "")
                D.add(
                    make_row(
                        account=company,
                        contact=get("contact"),
                        title=get("title"),
                        phone=get("phone"),
                        email=get("email"),
                        segment=force,
                        call_band="Call now" if "1" in bant_tier or "Call" in bant_tier else get("penetration") or "",
                        tier=bant_tier,
                        county=get("county"),
                        penetration=get("penetration"),
                        source=f"bant:{sheet}",
                        s_tier="S" if bant_tier.strip() in ("1", "Tier 1", "T1", "A") or str(bant_tier).startswith("1") else "",
                    )
                )
        wb.close()

    # 9) Metro Jax All_Orgs — full TAM coverage
    metro = Path("/workspace/b2b-jax/Metro-Jax-B2B-Segments-Orgs.xlsx")
    if metro.exists():
        wb = openpyxl.load_workbook(metro, read_only=True, data_only=True)
        ws = wb["All_Orgs"]
        rows_iter = ws.iter_rows(values_only=True)
        hdr = next(rows_iter)
        hdr = [str(h or "").strip() for h in hdr]
        idx = {h: i for i, h in enumerate(hdr)}

        def g(r, n):
            i = idx.get(n)
            return r[i] if i is not None and i < len(r) else None

        for r in rows_iter:
            if not r or not g(r, "OrgName"):
                continue
            call = str(g(r, "CallBand") or "")
            D.add(
                make_row(
                    account=g(r, "OrgName"),
                    contact_first=g(r, "ContactFirst"),
                    contact_last=g(r, "ContactLast"),
                    title=g(r, "Persona_Title") or g(r, "ICP"),
                    phone=g(r, "Phone"),
                    email=g(r, "Email"),
                    segment=g(r, "Segment"),
                    call_band=call,
                    county=g(r, "County"),
                    city=g(r, "City"),
                    notes=g(r, "Notes"),
                    penetration=g(r, "BuyerMotion"),
                    source="metro_jax:All_Orgs",
                    tier="Tier 1" if "call now" in call.lower() else "",
                )
            )
        wb.close()

    # ---------- Sort & KPIs ----------
    def sort_key(r: dict):
        seg = r.get("segment") or "Other"
        try:
            seg_i = SEGMENT_ORDER.index(seg)
        except ValueError:
            seg_i = 999
        return (
            seg_i,
            0 if r.get("is_tier1") else 1,
            0 if r.get("has_email") else 1,
            0 if r.get("has_phone") else 1,
            (r.get("account") or "").lower(),
            (r.get("contact") or "").lower(),
        )


    rows = sorted(D.rows, key=sort_key)

    # Ensure all SEGMENT_ORDER appear in chips even if zero
    seg_counts = Counter(r["segment"] for r in rows)
    tier1_by_seg = Counter(r["segment"] for r in rows if r.get("is_tier1"))
    email_by_seg = Counter(r["segment"] for r in rows if r.get("has_email"))

    segments_meta = []
    for s in SEGMENT_ORDER:
        segments_meta.append(
            {
                "id": s,
                "contacts": seg_counts.get(s, 0),
                "tier1": tier1_by_seg.get(s, 0),
                "with_email": email_by_seg.get(s, 0),
            }
        )
    # any extra segments not in order
    for s, n in seg_counts.most_common():
        if s not in SEGMENT_ORDER:
            segments_meta.append(
                {"id": s, "contacts": n, "tier1": tier1_by_seg.get(s, 0), "with_email": email_by_seg.get(s, 0)}
            )

    accounts = {r["account"].lower() for r in rows if r.get("account")}
    kpis = {
        "segments": len([s for s in segments_meta if s["contacts"] > 0]),
        "segments_listed": len(segments_meta),
        "accounts": len(accounts),
        "contacts": len(rows),
        "with_email": sum(1 for r in rows if r.get("has_email")),
        "with_phone": sum(1 for r in rows if r.get("has_phone")),
        "with_named_contact": sum(1 for r in rows if r.get("contact")),
        "tier1": sum(1 for r in rows if r.get("is_tier1")),
        "tier1_with_email": sum(1 for r in rows if r.get("is_tier1") and r.get("has_email")),
        "has_info": sum(1 for r in rows if r.get("has_info")),
        "gaps_zero_contacts": [s["id"] for s in segments_meta if s["contacts"] == 0],
    }

    payload = {
        "as_of": AS_OF,
        "product": "Sheehan Homestead / Critters on Call — Unified BDR TAM Dialer",
        "rule": "Real emails/phones only — never invented. Sort: Tier1 → email → phone → name within each sector.",
        "dedupe": "email → phone → org+name (merge richer fields)",
        "kpis": kpis,
        "segments": segments_meta,
        "sources": [{"id": k, "rows_ingested": v} for k, v in D.source_counts.most_common()],
        "contacts": rows,
    }

    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=None, separators=(",", ":")))
    print(json.dumps({"wrote": str(OUT), "bytes": OUT.stat().st_size, "kpis": kpis}, indent=2))
    print("gaps:", kpis["gaps_zero_contacts"])
    print("top segments:", [(s["id"], s["contacts"], s["tier1"], s["with_email"]) for s in segments_meta[:15]])


if __name__ == "__main__":
    main()
